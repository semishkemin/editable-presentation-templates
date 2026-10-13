"""Create a colorful timeline using editable PowerPoint shapes.

Requires Python 3.10+ and python-pptx 1.0.2.
Install: python -m pip install python-pptx==1.0.2
Run:     python 01_editable_timeline.py
Single:  python 01_editable_timeline.py --single
Custom:  python 01_editable_timeline.py --input templates/01_editable_timeline/example.json

The default output contains a cover and four examples with 3 to 6 milestones.
Edit ITEMS and SETTINGS below, or provide a JSON file with the same fields.
Layout recalculates when this script runs. PowerPoint edits do not trigger it.
Milestones use equal spacing, not a proportional calendar scale.
"""

from __future__ import annotations

import argparse
import json
import re
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


# 1. CONTENT: Keep six items for the default pack. Use --single for 3 to 6 items.
ITEMS = [
    {"label": "WEEK 01", "title": "Explore", "description": "Identify the topic and gather initial questions."},
    {"label": "WEEK 02", "title": "Plan", "description": "Set a clear goal and choose the next steps."},
    {"label": "WEEK 03", "title": "Create", "description": "Build a first version and make the ideas visible."},
    {"label": "WEEK 04", "title": "Review", "description": "Collect feedback and identify useful changes."},
    {"label": "WEEK 05", "title": "Improve", "description": "Refine the work using what you have learned."},
    {"label": "WEEK 06", "title": "Share", "description": "Present the result and record the next steps."},
]

# 2. APPEARANCE: Colors use six hexadecimal digits without a leading hash.
SETTINGS = {
    "title": "Project timeline",
    "subtitle": "A sequence of milestones",
    "font": "Arial",
    "background": "F6F3ED",
    "text": "20283D",
    "muted": "596274",
    "palette": ["5B4BDB", "007F82", "BB4B35", "315ED1", "A53575", "687326"],
}

SUBSTACK_URL = "https://substack.com/@semishkemin"
PROJECT_ROOT = Path(__file__).resolve().parent

SLIDE_WIDTH = 13.333333
SLIDE_HEIGHT = 7.5


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def tint(value: str, amount: float = 0.90) -> str:
    """Blend a color toward white for a light, opaque card background."""
    channels = [int(value[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in channels)


def wrapped(value: str, width: int, max_lines: int, field: str) -> str:
    """Use conservative wrapping and reject text that exceeds the template."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must be a non-empty string.")
    lines = textwrap.wrap(" ".join(value.split()), width=width)
    if len(lines) > max_lines:
        raise ValueError(f"{field} is too long. Use at most {max_lines} short lines.")
    return "\n".join(lines)


def validate(items: list, settings: dict) -> None:
    if not isinstance(items, list) or not 3 <= len(items) <= 6:
        raise ValueError("Provide between 3 and 6 timeline items.")
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            raise ValueError(f"Item {index} must be an object.")
        for key in ("label", "title", "description"):
            if key not in item:
                raise ValueError(f"Item {index} is missing '{key}'.")
        wrapped(item["label"], 20, 1, f"Item {index} label")
        wrapped(item["title"], 21, 2, f"Item {index} title")
        wrapped(item["description"], 24, 3, f"Item {index} description")
    wrapped(settings["title"], 45, 1, "Slide title")
    wrapped(settings["subtitle"], 85, 1, "Slide subtitle")
    if not isinstance(settings["font"], str) or not settings["font"].strip():
        raise ValueError("Font must be a non-empty string.")
    palette = settings["palette"]
    if not isinstance(palette, list) or not palette:
        raise ValueError("Palette must be a non-empty list of colors.")
    for color in [settings["background"], settings["text"], settings["muted"], *palette]:
        if not isinstance(color, str) or not re.fullmatch(r"[0-9a-fA-F]{6}", color):
            raise ValueError(f"Invalid color: {color!r}. Use six hexadecimal digits.")


def shape(slide, kind, x, y, width, height, fill, name):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    obj.name = name
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background()
    return obj


def text(slide, value, x, y, width, height, size, color, font, name,
         bold=False, centered=False):
    """Keep every label as real, editable text rather than vector outlines."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if centered else MSO_ANCHOR.TOP
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.1
        p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    return box


def add_timeline(prs, items, settings):
    """Lay out alternating cards around a segmented horizontal time band."""
    validate(items, settings)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(settings["background"])
    font = settings["font"]
    text(slide, settings["title"], 0.65, 0.42, 12.0, 0.62, 32,
         settings["text"], font, "Slide title", bold=True)
    text(slide, settings["subtitle"], 0.68, 1.12, 11.95, 0.38, 17,
         settings["muted"], font, "Slide subtitle")

    left, right, axis_y = 0.70, 12.63, 4.22
    cell_width = (right - left) / len(items)
    card_width, card_height = min(3.48, cell_width * 1.50), 2.30

    # Draw the time band first so milestone circles sit above it.
    for i in range(len(items)):
        color = settings["palette"][i % len(settings["palette"])]
        shape(slide, MSO_SHAPE.CHEVRON, left + i * cell_width, axis_y - 0.16,
              cell_width + (0.10 if i < len(items) - 1 else 0), 0.32,
              color, f"Time band {i + 1}")

    for i, item in enumerate(items):
        number = i + 1
        color = settings["palette"][i % len(settings["palette"])]
        center_x = left + cell_width * (i + 0.5)
        card_x = max(left, min(center_x - card_width / 2, right - card_width))
        above = i % 2 == 0
        card_y = 1.65 if above else 4.74
        stem_top = card_y + card_height if above else axis_y
        stem_bottom = axis_y if above else card_y
        shape(slide, MSO_SHAPE.RECTANGLE, center_x - 0.022, stem_top,
              0.044, stem_bottom - stem_top, color, f"Milestone {number} stem")
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y,
              card_width, card_height, tint(color), f"Milestone {number} card")
        shape(slide, MSO_SHAPE.RECTANGLE, card_x + 0.20, card_y + 0.22,
              0.055, 0.25, color, f"Milestone {number} accent")
        text(slide, item["label"], card_x + 0.36, card_y + 0.20,
             card_width - 0.55, 0.30, 13, color, font, f"Milestone {number} label", bold=True)
        text(slide, wrapped(item["title"], 21, 2, "Title"), card_x + 0.20,
             card_y + 0.59, card_width - 0.40, 0.65, 20, settings["text"],
             font, f"Milestone {number} title", bold=True)
        text(slide, wrapped(item["description"], 24, 3, "Description"), card_x + 0.20,
             card_y + 1.21, card_width - 0.40, 0.98, 17, settings["muted"],
             font, f"Milestone {number} description")
        shape(slide, MSO_SHAPE.OVAL, center_x - 0.29, axis_y - 0.29,
              0.58, 0.58, settings["background"], f"Milestone {number} halo")
        shape(slide, MSO_SHAPE.OVAL, center_x - 0.235, axis_y - 0.235,
              0.47, 0.47, color, f"Milestone {number} marker")
        text(slide, str(number), center_x - 0.235, axis_y - 0.235, 0.47, 0.47,
             15, "FFFFFF", font, f"Milestone {number} number", bold=True, centered=True)
    return slide


def add_cover(prs, settings):
    """Add a removable cover with usage permissions and a clickable profile."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(settings["background"])
    font, ink = settings["font"], settings["text"]
    accent = settings["palette"][0]
    text(slide, "Editable timelines", 0.75, 0.72, 11.8, 0.85, 44,
         ink, font, "Cover title", bold=True)
    text(slide, "For lessons, workshops, and presentations", 0.78, 1.73,
         11.7, 0.55, 23, settings["muted"], font, "Cover subtitle")
    text(slide, "Make these templates your own", 0.78, 2.75, 11.7, 0.55,
         25, accent, font, "Usage heading", bold=True)
    text(slide, "Change the text, colors, fonts, and shapes.\n"
         "Use, adapt, and redistribute these templates freely.\n"
         "Personal, educational, and commercial use is welcome.\n"
         "No attribution is required. You may remove this cover.",
         0.78, 3.47, 11.65, 1.75, 21, ink, font, "Usage permissions")
    text(slide, "Created for my own presentations and shared with you.",
         0.78, 5.62, 11.7, 0.40, 18, settings["muted"], font, "Creator note")
    link = text(slide, "substack.com/@semishkemin", 0.78, 6.13, 11.7, 0.43,
                21, accent, font, "Substack profile", bold=True)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = SUBSTACK_URL
    text(slide, "Feedback is welcome: what worked, and what would you change?",
         0.78, 6.77, 11.7, 0.35, 16, settings["muted"], font, "Feedback invitation")
    return slide


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--single", action="store_true", help="Create one timeline using ITEMS, plus a cover.")
    mode.add_argument("--input", type=Path, help="Create one timeline from a UTF-8 JSON file, plus a cover.")
    mode.add_argument("--demo", action="store_true", help="Alias for the default four-example pack.")
    parser.add_argument("--no-cover", action="store_true", help="Omit the cover slide.")
    parser.add_argument("--output", type=Path, help="Destination PPTX file. Existing files are replaced.")
    args = parser.parse_args()
    settings, items = dict(SETTINGS), ITEMS
    if args.input:
        data = json.loads(args.input.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or "items" not in data:
            raise ValueError("The JSON file must contain an 'items' list.")
        overrides = data.get("settings", {})
        if not isinstance(overrides, dict) or set(overrides) - set(SETTINGS):
            raise ValueError("Settings must be an object containing only supported settings.")
        settings.update(overrides)
        items = data["items"]
    validate(items, settings)
    pack = not (args.single or args.input)
    if pack and len(items) != 6:
        raise ValueError("The example pack needs exactly 6 ITEMS. Use --single for fewer items.")
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH)
    prs.slide_height = Inches(SLIDE_HEIGHT)
    if not args.no_cover:
        add_cover(prs, settings)
    if pack:
        for count in range(3, 7):
            add_timeline(prs, items[:count], {**settings, "title": f"{settings['title']}: {count} milestones"})
    else:
        add_timeline(prs, items, settings)
    name = "timeline_examples.pptx" if pack else "custom_timeline.pptx"
    output = args.output or PROJECT_ROOT / "exports" / "01_editable_timeline" / name
    if output.suffix.lower() != ".pptx":
        raise ValueError("The output filename must end with .pptx.")
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Saved {len(prs.slides)} editable slide(s) to {output.resolve()}")
    print("PNG previews are supplied separately. Regenerate them after editing the deck.")


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError) as error:
        raise SystemExit(f"Error: {error}") from error
